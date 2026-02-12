from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook
from orders.models import Order, PickUpPoint
from products.models import Product
from pathlib import Path
from datetime import datetime

class Command(BaseCommand):
    help = 'Импорт заказов (сокращенно, пропускает ошибки)'

    def handle(self, *args, **options):
        path = Path('import/Заказ_import.xlsx')
        if not path.exists(): return

        ws = load_workbook(path).active
        count = 0

        with transaction.atomic():
            for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                if not any(row): continue
                (num, arts, d1, d2, p_id, name, code, stat) = row[:8]

                try:
                    # Преобразование дат: если ошибка (н-р 30.02), уйдет в общий except и пропустит строку
                    if not isinstance(d1, datetime): d1 = datetime.strptime(str(d1).strip('“” '), '%d.%m.%Y')
                    if not isinstance(d2, datetime): d2 = datetime.strptime(str(d2).strip('“” '), '%d.%m.%Y')

                    # Поиск товара
                    art = str(arts).split(',')[0].strip()
                    product = Product.objects.filter(article=art).first() or \
                              Product.objects.filter(article=art.replace('А', 'A')).first()
                    
                    if not product: # Если товар не найден даже с заменой 'А'
                        self.stdout.write(self.style.WARNING(f"[{idx}] Товар '{art}' не найден. Пропускаю строку."))
                        continue 

                    # Поиск пункта выдачи
                    point = PickUpPoint.objects.get(id=int(p_id))

                    # Создание/Обновление заказа
                    Order.objects.update_or_create(
                        order_number=num,
                        defaults={
                            'product': product, 'order_date': d1, 'delivery_date': d2,
                            'pick_up_point': point,
                            'customer_full_name': name, 'receive_code': code, 'status': stat or 'Новый'
                        }
                    )
                    count += 1
                    # self.stdout.write(f"[{idx}] Обработан заказ {num}") # Опциональный лог успеха
                except Exception as e:
                    # Логирование ошибки и пропуск строки
                    self.stdout.write(self.style.ERROR(f"[{idx}] Ошибка обработки заказа {num}: {e}. Пропускаю строку."))
                    continue 

        self.stdout.write(self.style.SUCCESS(f'Готово. Успешно импортировано/обновлено заказов: {count}'))
