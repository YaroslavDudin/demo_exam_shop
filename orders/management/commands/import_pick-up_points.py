from django.core.management.base import BaseCommand
from openpyxl import load_workbook
from orders.models import PickUpPoint
from django.db import models  


class Command(BaseCommand):
    help = 'Импорт пунктов выдачи из Excel'

    def handle(self, *args, **kwargs):
        wb = load_workbook(r'import\Пункты выдачи_import.xlsx').active

        # Определяем стартовый номер
        last_number = PickUpPoint.objects.aggregate(models.Max('number'))['number__max'] or 0
        count = 0

        for row in wb.iter_rows(min_row=1, values_only=True):
            address = row[0]

            if address:
                last_number += 1
                PickUpPoint.objects.create(address=address, number=last_number)
                count += 1

        self.stdout.write(self.style.SUCCESS(
            f'Импортировано пунктов выдачи: {count}'
        ))
