from django.core.management.base import BaseCommand
from products.models import Product, Category, Manufacturer, Supplier
from openpyxl import load_workbook
from pathlib import Path

class Command(BaseCommand):
    help = 'Импорт товаров из Excel (упрощённо)'

    def handle(self, *args, **options):
        wb = load_workbook('import/Tovar.xlsx').active
        media_dir = Path('media/products')

        for idx, row in enumerate(wb.iter_rows(min_row=2, values_only=True), start=2):
            article, name, unit, price, supplier_name, manufacturer_name, category_name, discount, quantity, description, image_name = row
            if not article:
                continue

            category, _ = Category.objects.get_or_create(name=category_name)
            manufacturer, _ = Manufacturer.objects.get_or_create(name=manufacturer_name)
            supplier, _ = Supplier.objects.get_or_create(name=supplier_name)

            product, _ = Product.objects.get_or_create(
                article=article,
                defaults={
                    'name': name,
                    'unit': unit ,
                    'price': price ,
                    'discount': discount ,
                    'quantity': quantity ,
                    'description': description,
                    'category': category,
                    'manufacturer': manufacturer,
                    'supplier': supplier
                }
            )

            if image_name:
                image_path = media_dir / str(image_name).strip()
                if image_path.exists() and not product.image.name:
                    product.image.name = f'products/{image_path.name}'
                    product.save(update_fields=['image'])

            self.stdout.write(f"[{idx}] Импортирован: {article}")

        self.stdout.write(self.style.SUCCESS("Импорт товаров завершён"))
